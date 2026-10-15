"""DOCX -> plain text. Gemini cannot read .docx directly (SPEC §4)."""

import io

from docx import Document


def docx_to_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())
